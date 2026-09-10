# -*- coding: utf-8 -*-
"""Устойчивость разбора к изменению шаблона.

Главная проверка проекта: рапорт собирают вручную, и его верстка живёт
своей жизнью. Разбор не должен зависеть от номеров строк и столбцов —
здесь это подтверждается искусственным сдвигом листа и переименованием.
"""
from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from rvp_report.excel.resolve import resolve
from rvp_report.excel.sheet import Sheet

from .conftest import report_files

#: Сдвиги «строки, столбцы», которые разбор обязан пережить.
SHIFTS = ((7, 0), (100, 0), (0, 3), (12, 4))


@pytest.fixture(scope="module")
def sample() -> Path:
    """Один рапорт, на котором проверяется устойчивость."""
    files = report_files()
    if not files:
        pytest.skip("нет рапортов для проверки")
    return files[0]


def _sheet(path: Path) -> Sheet:
    """Лист рапорта как есть."""
    return Sheet(openpyxl.load_workbook(path, data_only=True)["Сут.рапорт"])


def _shifted(path: Path, dr: int, dc: int) -> Sheet:
    """Копия листа, сдвинутая на dr строк и dc столбцов, как это делает Excel.

    Переносятся и значения, и объединённые ячейки. Это существенно:
    ``openpyxl.insert_rows`` объединения не двигает, и проверка на нём
    получилась бы неправдоподобно лёгкой.
    """
    source = openpyxl.load_workbook(path, data_only=True)["Сут.рапорт"]
    target = openpyxl.Workbook().active
    target.title = "Сут.рапорт"
    for row in source.iter_rows(min_row=1, max_row=source.max_row,
                                max_col=source.max_column):
        for cell in row:
            if cell.value is not None:
                target.cell(row=cell.row + dr, column=cell.column + dc,
                            value=cell.value)
    for merged in source.merged_cells.ranges:
        target.merge_cells(start_row=merged.min_row + dr,
                           start_column=merged.min_col + dc,
                           end_row=merged.max_row + dr,
                           end_column=merged.max_col + dc)
    return Sheet(target)


@pytest.mark.parametrize("dr,dc", SHIFTS, ids=[f"+{r}стр+{c}стб"
                                               for r, c in SHIFTS])
def test_blocks_follow_content(sample: Path, dr: int, dc: int) -> None:
    """Разделы и столбцы должны поехать ровно за содержимым."""
    original = resolve(_sheet(sample))
    shifted = resolve(_shifted(sample, dr, dc))

    assert not shifted.degraded, (
        f"сдвиг +{dr} строк, +{dc} столбцов: " + "; ".join(shifted.degraded))

    for key, (start, _end) in original.blocks.items():
        moved = shifted.blocks.get(key)
        assert moved is not None, f"потерян раздел «{key}»"
        assert moved[0] == start + dr, (
            f"раздел «{key}» ожидался в строке {start + dr}, "
            f"найден в {moved[0]}")

    for table, columns in original.columns.items():
        expected = {name: col + dc for name, col in columns.items()}
        assert shifted.columns.get(table) == expected, (
            f"разъехались столбцы таблицы «{table}»")


def test_renamed_sheet_is_recognised(sample: Path) -> None:
    """Переименованный лист должен опознаваться по содержимому."""
    workbook = openpyxl.load_workbook(sample, data_only=True)
    worksheet = workbook["Сут.рапорт"]
    worksheet.title = "Лист 1"
    resolved = resolve(Sheet(worksheet))
    assert "stock" in resolved.blocks, (
        "на переименованном листе не найден раздел фонда скважин")
