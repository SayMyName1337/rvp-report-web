# -*- coding: utf-8 -*-
"""Проверки редактора разметки.

Главное, что здесь проверяется, — не то, что поправка работает, а то,
что она **не может навредить**: на неизменённом файле она не срабатывает
вовсе, а прежние названия продолжают находиться.
"""
from __future__ import annotations

import shutil
from pathlib import Path

import openpyxl
import pytest
from django.urls import reverse

from apps.accounts.models import Role, User, ensure_roles
from apps.layout import services
from apps.layout.models import LayoutFix, active_fixes
from apps.reports.loader import read_file, store
from apps.reports.models import LayoutGap, Report
from rvp_report.excel.resolve import resolve
from rvp_report.excel.sheet import Sheet
from rvp_report.excel.spec import Fix

from .conftest import report_files

#: Строка и столбец подписи, по которой находится прогноз сдачи.
LABEL_ROW, LABEL_COL = 564, 2
#: Как ЦИТУ мог бы её переименовать.
NEW_LABEL = "Итого по отгрузке"


@pytest.fixture
def sample() -> Path:
    """Настоящий суточный рапорт."""
    files = report_files()
    if not files:
        pytest.skip("нет рапортов для проверки")
    return files[0]


@pytest.fixture
def renamed(sample: Path, tmp_path: Path) -> Path:
    """Копия рапорта, в которой подпись строки переименована."""
    target = tmp_path / sample.name
    shutil.copy(sample, target)
    workbook = openpyxl.load_workbook(target)
    workbook["Сут.рапорт"].cell(row=LABEL_ROW, column=LABEL_COL).value = \
        NEW_LABEL
    workbook.save(target)
    return target


def _sheet(path: Path) -> Sheet:
    """Лист рапорта."""
    return Sheet(openpyxl.load_workbook(path, data_only=True)["Сут.рапорт"])


# --- движок ------------------------------------------------------------------
def test_rename_is_noticed(renamed: Path) -> None:
    """Переименованная подпись должна дать разобранный пробел разметки."""
    resolved = resolve(_sheet(renamed))
    keys = {gap.key for gap in resolved.gaps}
    assert "forecast_total" in keys, "переименование осталось незамеченным"
    gap = next(g for g in resolved.gaps if g.key == "forecast_total")
    assert gap.label == "Всего"
    assert gap.reference, "не записан эталонный адрес"


def test_fix_repairs_the_rename(renamed: Path) -> None:
    """Поправка возвращает показатель на место."""
    fixes = {"forecast_total": Fix(row_aliases=(NEW_LABEL,))}
    resolved = resolve(_sheet(renamed), fixes)
    assert not [g for g in resolved.gaps if g.key == "forecast_total"]
    assert resolved.applied, "не отмечено, что сработала поправка"


def test_fix_does_not_fire_on_unchanged_file(sample: Path) -> None:
    """На неизменённом файле поправка не должна срабатывать вовсе.

    Это главное свойство механизма: синоним добавляется к подписям схемы,
    а не заменяет их, поэтому испортить работавшее он не может.
    """
    fixes = {"forecast_total": Fix(row_aliases=(NEW_LABEL,))}
    plain = resolve(_sheet(sample))
    with_fix = resolve(_sheet(sample), fixes)

    assert with_fix.applied == [], "поправка сработала там, где не нужна"
    assert with_fix.cells["forecast_total"] == plain.cells["forecast_total"]


def test_old_name_keeps_working(renamed: Path, sample: Path) -> None:
    """Прежнее название должно продолжать находиться.

    Иначе рапорт за прошлый месяц перестанет разбираться после того,
    как подпись поменяли.
    """
    fixes = {"forecast_total": Fix(row_aliases=("На ПСП Головные", NEW_LABEL))}
    for path in (sample, renamed):
        resolved = resolve(_sheet(path), fixes)
        assert not [g for g in resolved.gaps if g.key == "forecast_total"]


# --- вывод правила -----------------------------------------------------------
def test_rule_prefers_label_over_address(sample: Path) -> None:
    """Из клетки с подписью слева выводится название, а не адрес.

    Подпись переживает вставку строк, адрес — нет.
    """
    sheet = _sheet(sample)
    kind, value = services.derive_rule(sheet, LABEL_ROW, 4, "row")
    assert kind == LayoutFix.Kind.ROW
    assert value == "На ПСП Головные"


def test_rule_falls_back_to_address(sample: Path) -> None:
    """Там, где подписи нет, остаётся привязка к адресу."""
    sheet = _sheet(sample)
    kind, value = services.derive_rule(sheet, 1, 1, "row")
    assert kind == LayoutFix.Kind.ADDRESS
    assert value == "A1"


def test_candidate_points_at_the_renamed_row(renamed: Path) -> None:
    """Догадка должна указать на строку, стоящую там же, где прежняя."""
    sheet = _sheet(renamed)
    found = services.candidates(sheet, "Всего", (LABEL_ROW, 4),
                                (LABEL_ROW - 10, LABEL_ROW + 10))
    assert found, "не предложено ни одной догадки"
    assert found[0].label == NEW_LABEL


# --- хранение поправок -------------------------------------------------------
@pytest.mark.django_db
def test_aliases_of_one_metric_are_merged() -> None:
    """Несколько поправок на показатель собираются в один набор названий."""
    user = User.objects.create_user(username="citu", password="x" * 14,
                                    must_change_password=False)
    for value in ("Первое имя", "Второе имя"):
        LayoutFix.objects.create(key="forecast_total", label="Всего",
                                 kind=LayoutFix.Kind.ROW, value=value,
                                 created_by=user)
    fixes = active_fixes()
    assert fixes["forecast_total"].row_aliases == ("Первое имя", "Второе имя")


@pytest.mark.django_db
def test_cancelled_fix_is_not_applied() -> None:
    """Отменённая поправка не должна попадать в разбор."""
    user = User.objects.create_user(username="citu", password="x" * 14,
                                    must_change_password=False)
    LayoutFix.objects.create(key="forecast_total", label="Всего",
                             kind=LayoutFix.Kind.ROW, value="Имя",
                             created_by=user, is_active=False)
    assert active_fixes() == {}


# --- загрузка ----------------------------------------------------------------
@pytest.mark.django_db
def test_gap_marks_report_and_disappears_after_fix(renamed: Path) -> None:
    """Пробел помечает сутки, а после поправки исчезает."""
    daily, gaps = read_file(renamed)
    report = store(daily, gaps)
    assert report.status == Report.Status.NEEDS_REVIEW
    assert LayoutGap.objects.filter(key="forecast_total").exists()

    # Одна подпись держала два показателя — так же, как их чинит
    # страница подтверждения: одним указанием сразу оба.
    user = User.objects.create_user(username="citu", password="x" * 14,
                                    must_change_password=False)
    for key in ("forecast_total", "forecast_rate_tph"):
        LayoutFix.objects.create(key=key, label="Всего",
                                 kind=LayoutFix.Kind.ROW, value=NEW_LABEL,
                                 created_by=user)

    daily, gaps = read_file(renamed)
    report = store(daily, gaps)
    assert report.status == Report.Status.LOADED
    assert not LayoutGap.objects.exists()
    assert report.delivery.forecast_total is not None
    assert report.delivery.forecast_rate_tph is not None


# --- доступ ------------------------------------------------------------------
@pytest.mark.django_db
def test_editor_is_closed_to_viewers(client) -> None:
    """Руководство не правит разметку: это дело работника промысла."""
    groups = ensure_roles()
    viewer = User.objects.create_user(username="chief", password="x" * 14,
                                      must_change_password=False)
    viewer.groups.add(groups[Role.MANAGEMENT.value])
    client.force_login(viewer)
    assert client.get(reverse("layout:pending")).status_code == 403


@pytest.mark.django_db
def test_editor_is_open_to_field_worker(client) -> None:
    """Работнику промысла редактор доступен."""
    groups = ensure_roles()
    operator = User.objects.create_user(username="citu", password="x" * 14,
                                        must_change_password=False)
    operator.groups.add(groups[Role.FIELD.value])
    client.force_login(operator)
    assert client.get(reverse("layout:pending")).status_code == 200
