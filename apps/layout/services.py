# -*- coding: utf-8 -*-
"""Работа редактора разметки: показать лист, предложить догадку, вывести правило.

Оператор не пишет правил. Он видит кусок таблицы, которую сам заполняет
каждый день, и щёлкает по нужной клетке — а правило из этого выводит
приложение. Здесь собрано всё, что для этого нужно.
"""
from __future__ import annotations

from dataclasses import dataclass
from difflib import SequenceMatcher
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter

from rvp_report.excel.sheet import Sheet, norm

from .models import LayoutFix

#: Сколько строк и столбцов показывать вокруг предполагаемого места.
RADIUS_ROWS = 6
RADIUS_COLS = 5

#: Насколько подписи должны быть похожи, чтобы предлагать их как догадку.
SIMILARITY = 0.55

#: Сколько ячеек влево просматривать в поисках подписи строки.
LABEL_REACH = 8


@dataclass(frozen=True)
class Cell:
    """Одна клетка показываемого куска листа."""

    row: int
    col: int
    text: str
    rowspan: int = 1
    colspan: int = 1
    is_number: bool = False
    is_reference: bool = False

    @property
    def address(self) -> str:
        """Адрес клетки в привычном виде, например «D564»."""
        return f"{get_column_letter(self.col)}{self.row}"


@dataclass(frozen=True)
class Candidate:
    """Догадка о том, где теперь лежит показатель."""

    row: int
    col: int
    label: str
    value: str
    why: str

    @property
    def address(self) -> str:
        """Адрес предполагаемой ячейки."""
        return f"{get_column_letter(self.col)}{self.row}"


def open_report(path: str | Path) -> Sheet:
    """Открыть лист суточного рапорта для показа.

    Args:
        path: Путь к сохранённому файлу рапорта.

    Returns:
        Лист с развёрнутыми объединениями.
    """
    workbook = openpyxl.load_workbook(path, data_only=True)
    name = "Сут.рапорт"
    worksheet = workbook[name] if name in workbook.sheetnames \
        else workbook[workbook.sheetnames[0]]
    return Sheet(worksheet)


def _text(value) -> str:
    """Значение ячейки в виде короткой строки."""
    if value is None:
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).replace("\n", " ").strip()


def fragment(sheet: Sheet, row: int, col: int,
             reference: tuple[int, int] | None = None) -> dict:
    """Кусок листа вокруг указанной клетки — как его видит оператор.

    Объединённые ячейки показываются одной клеткой: иначе таблица
    на экране не совпадёт с тем, что человек видит в Excel.

    Args:
        sheet: Лист рапорта.
        row: Строка, вокруг которой показывать.
        col: Столбец, вокруг которого показывать.
        reference: Ячейка, взятая по эталону, — будет подсвечена.

    Returns:
        Заголовки столбцов и строки клеток для отрисовки таблицы.
    """
    top = max(1, row - RADIUS_ROWS)
    bottom = min(sheet.max_row, row + RADIUS_ROWS)
    left = max(1, col - RADIUS_COLS)
    right = min(sheet.max_col, col + RADIUS_COLS)

    covered: set[tuple[int, int]] = set()
    rows: list[dict] = []
    for r in range(top, bottom + 1):
        cells: list[Cell] = []
        for c in range(left, right + 1):
            if (r, c) in covered:
                continue
            span = sheet.span(r, c)
            anchor_row, anchor_col = span.row, span.col
            for rr in span.rows:
                for cc in span.cols:
                    covered.add((rr, cc))
            value = sheet.raw(anchor_row, anchor_col)
            cells.append(Cell(
                row=anchor_row, col=anchor_col, text=_text(value),
                rowspan=min(span.last_row, bottom) - r + 1,
                colspan=min(span.last_col, right) - c + 1,
                is_number=isinstance(value, (int, float))
                and not isinstance(value, bool),
                is_reference=reference is not None
                and (anchor_row, anchor_col) == reference))
        rows.append({"number": r, "cells": cells})

    return {"rows": rows,
            "columns": [get_column_letter(c) for c in range(left, right + 1)],
            "top": top, "left": left}


def row_label(sheet: Sheet, row: int, col: int) -> str:
    """Подпись строки: ближайший текст слева от клетки.

    Args:
        sheet: Лист рапорта.
        row: Строка клетки.
        col: Столбец клетки.

    Returns:
        Текст подписи или пустую строку, если слева ничего нет.
    """
    span = sheet.span(row, col)
    for c in range(span.col - 1, max(0, span.col - LABEL_REACH), -1):
        value = sheet.value(row, c)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def column_header(sheet: Sheet, row: int, col: int, top: int) -> str:
    """Заголовок столбца: ближайший текст сверху в пределах блока.

    Args:
        sheet: Лист рапорта.
        row: Строка клетки.
        col: Столбец клетки.
        top: Верхняя граница блока — выше неё не поднимаемся.

    Returns:
        Текст заголовка или пустую строку.
    """
    for r in range(row - 1, max(top - 1, 0), -1):
        value = sheet.value(r, col)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return ""


def derive_rule(sheet: Sheet, row: int, col: int, part: str,
                block_top: int = 1) -> tuple[str, str]:
    """Вывести правило из клетки, на которую указал оператор.

    Опознание по подписи предпочтительнее привязки к адресу: подпись
    переживает вставку строк, адрес — нет. Поэтому адрес возвращается
    только тогда, когда ни подписи строки, ни заголовка столбца нет.

    Args:
        sheet: Лист рапорта.
        row: Строка выбранной клетки.
        col: Столбец выбранной клетки.
        part: Что искали — ``row`` или ``col``.
        block_top: Верхняя граница блока для поиска заголовка столбца.

    Returns:
        Пару «вид правила, значение» в терминах :class:`LayoutFix.Kind`.
    """
    if part == "col":
        header = column_header(sheet, row, col, block_top)
        if header:
            return LayoutFix.Kind.COL, header
    else:
        label = row_label(sheet, row, col)
        if label:
            return LayoutFix.Kind.ROW, label

    return LayoutFix.Kind.ADDRESS, f"{get_column_letter(col)}{row}"


def candidates(sheet: Sheet, label: str, reference: tuple[int, int] | None,
               block: tuple[int, int] | None = None) -> list[Candidate]:
    """Догадки о том, где теперь лежит показатель.

    Два источника, оба ненадёжные по отдельности и потому показываемые
    вместе: клетка, стоящая там же, где раньше, и подписи, похожие
    на искомую. Ничего не применяется само — решает оператор.

    Args:
        sheet: Лист рапорта.
        label: Подпись, которую искали.
        reference: Эталонный адрес, взятый взамен.
        block: Границы раздела, в котором шёл поиск.

    Returns:
        Догадки, от более правдоподобных к менее.
    """
    found: list[Candidate] = []
    seen: set[tuple[int, int]] = set()

    if reference is not None:
        row, col = reference
        near = row_label(sheet, row, col)
        if near:
            found.append(Candidate(
                row=row, col=col, label=near,
                value=_text(sheet.value(row, col)),
                why=f"стоит там же, где раньше «{label}»"))
            seen.add((row, col))

    # Одинаковые подписи вроде «Всего» есть в каждом втором блоке, поэтому
    # поиск похожих ведётся рядом с прежним местом, а не по всему листу:
    # «Всего» из блока ППД к прогнозу сдачи отношения не имеет.
    target = norm(label)
    top, bottom = block if block else (1, sheet.max_row)
    for r in range(top, min(bottom, sheet.max_row) + 1):
        for c in range(1, min(sheet.max_col, 30) + 1):
            value = sheet.raw(r, c)
            if not isinstance(value, str) or not value.strip():
                continue
            score = SequenceMatcher(None, target, norm(value)).ratio()
            if score < SIMILARITY or (r, c) in seen:
                continue
            numbers = sheet.numbers_right(r, sheet.span(r, c).last_col,
                                          limit=1, within=6)
            if not numbers:
                continue
            seen.add((r, c))
            found.append(Candidate(
                row=r, col=c, label=value.strip(),
                value=_text(numbers[0]),
                why=f"подпись похожа на «{label}»"))

    if reference is not None:
        anchor = reference[0]
        found.sort(key=lambda item: abs(item.row - anchor))
    return found[:5]
