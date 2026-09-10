# -*- coding: utf-8 -*-
"""Чтение суточного рапорта из .xlsx в модель DailyReport.

Читатель не знает адресов ячеек: он спрашивает их у resolve.py, который
находит блоки и столбцы по подписям. Если подпись не нашлась, используется
эталонный адрес, и об этом появляется запись в списке замечаний.
"""
from __future__ import annotations

import re
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from . import layout as L
from ..domain.checks import validate, ru as _ru
from ..domain.model import (
    DailyReport, Delivery, DrillingRow, GasRow, HSE, InjectionRow, PiggingRow,
    PowerState, Production, QualityRow, RampUpRow, StockRow, TankRow, Totals,
    WellRow,
)
from .resolve import resolve
from .sheet import Sheet, norm
from .spec import Resolved, Table

EMPTY = {"", "-", "—", "0", "нет", "н/д", "none"}


def _num(value) -> float | None:
    """Число или None. Строки вида «1 234,5» и «<4» тоже разбирает."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("\xa0", "").replace(" ", "")
    if not text or text.lower() in EMPTY:
        return None
    text = text.lstrip("<>≤≥~").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return None


def _text(value) -> str:
    """Значение ячейки в виде строки; прочерки считаются пустотой."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y")
    text = str(value).strip()
    return "" if text in {"-", "—"} else text


def _date(value) -> date | None:
    """Дата из datetime или из строки «дд.мм.гггг» — встречаются оба вида."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    match = re.match(r"^(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})$", _text(value))
    if match:
        day, month, year = (int(g) for g in match.groups())
        try:
            return date(year, month, day)
        except ValueError:
            return None
    return None


def _is_empty(text: str) -> bool:
    """Считается ли текст незаполненным."""
    return text.strip().lower() in EMPTY


def open_sheet(path: str) -> Sheet:
    """Лист рапорта: по имени, а если переименовали — по содержимому."""
    wb = openpyxl.load_workbook(path, data_only=True)
    for name in L.SHEET_CANDIDATES:
        if name in wb.sheetnames:
            return Sheet(wb[name])

    best, best_score = None, 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        sheet = Sheet(ws)
        score = sum(1 for m in L.SHEET_MARKERS if sheet.has_label(m))
        if score > best_score:
            best, best_score = sheet, score
    if best is not None and best_score:
        return best
    raise ValueError(
        "Не удалось определить лист суточного рапорта. "
        f"Листы в файле: {', '.join(wb.sheetnames)}")


class ReportReader:
    """Чтение суточного рапорта по разрешённой схеме.

    Порядок работы: открыть лист, привязать схему, прочитать блоки,
    сверить величины между собой и записать происхождение показателей.
    """
    def __init__(self, path: str):
        """Args:
            path: Путь к файлу суточного рапорта .xlsx.

        Raises:
            ValueError: Если лист рапорта не удалось опознать.
        """
        self.path = path
        self.sheet = open_sheet(path)
        self.map: Resolved = resolve(self.sheet)
        self.report = DailyReport()

    # --- доступ через разрешённую схему --------------------------------------
    def _at(self, key: str):
        """Значение показателя по ключу разрешённой схемы."""
        cell = self.map.cell(key)
        return None if cell is None else self.sheet.value(*cell)

    def _num_at(self, key: str) -> float | None:
        """Числовое значение показателя или None."""
        return _num(self._at(key))

    def _text_at(self, key: str) -> str:
        """Текстовое значение показателя."""
        return _text(self._at(key))

    def _cell(self, table: str, column: str, row: int):
        """Значение ячейки таблицы на пересечении строки и столбца."""
        col = self.map.col(table, column)
        return None if col is None else self.sheet.value(row, col)

    def _rows_of(self, table: str, spec: Table):
        """Строки данных таблицы, уже отфильтрованные от подытогов и пустых."""
        rows = self.map.rows.get(table) or spec.reference_rows
        for row in range(rows[0], rows[1] + 1):
            key = _text(self._cell(table, spec.key_column, row))
            if _is_empty(key):
                continue
            if any(_is_empty(_text(self._cell(table, need, row)))
                   for need in spec.require):
                continue
            # строки подытогов «затекают» из объединённых ячеек: там, где
            # должно быть месторождение, оказывается число
            if any(_num(self._cell(table, need, row)) is not None
                   for need in spec.require_text):
                continue
            if any(_num(self._cell(table, need, row)) is None
                   for need in spec.require_number):
                continue
            yield row, key

    # --- разбор --------------------------------------------------------------
    def read(self) -> DailyReport:
        """Прочитать рапорт целиком.

        Returns:
            Заполненную модель со списком замечаний и журналом
            происхождения показателей.
        """
        self._carry_layout_notes()
        self._read_header()
        self._read_production()
        self._read_delivery()
        self._read_stock()
        self._read_movement()
        self._read_wells("idle", L.IDLE, "Бездействие")
        self._read_wells("downtime", L.DOWNTIME, "Текущий простой")
        self._read_vsp()
        self._read_ramp_up()
        self._read_injection()
        self._read_tanks()
        self._read_quality()
        self._read_gas()
        self._read_power()
        self._read_pigging()
        self._read_drilling()
        self._read_hse()
        validate(self.report)
        self._collect_sources()
        return self.report

    def _carry_layout_notes(self) -> None:
        """Перенести замечания о разметке в список замечаний рапорта."""
        for message in self.map.degraded:
            self.report.add_issue("warning", "Разметка файла", message)
        for message in self.map.notes:
            self.report.add_issue("info", "Разметка файла", message)

    def _read_header(self) -> None:
        """Дата, температура и ветер из шапки рапорта."""
        r = self.report
        r.report_date = _date(self._at("date"))
        r.temperature = self._num_at("temperature")
        r.wind = self._text_at("wind")
        if r.report_date is None:
            r.add_issue("error", "Шапка", "Не удалось прочитать дату рапорта")

    def _read_production(self) -> None:
        """Показатели добычи и распределение по месторождениям."""
        p = Production()
        for name in L.PRODUCTION:
            if hasattr(p, name):
                setattr(p, name, self._num_at(name))
        p.by_field = self._read_fields()
        self.report.production = p

    def _read_fields(self) -> list[tuple[str, float]]:
        """Добыча по месторождениям: столбцы опознаются по сокращениям в шапке."""
        cell = self.map.cell("fields_header")
        if cell is None:
            return []
        header_row = cell[0]
        value_row = header_row + L.FIELDS_VALUE_OFFSET
        known = {norm(short): full for short, full in L.FIELD_NAMES.items()}

        found: list[tuple[int, str]] = []
        unknown: list[tuple[int, str]] = []
        for col in range(1, self.sheet.max_col + 1):
            label = self.sheet.raw(header_row, col)
            if not isinstance(label, str) or not label.strip():
                continue
            key = norm(label)
            if key in known:
                found.append((col, known[key]))
            elif len(key) <= 12:
                unknown.append((col, label.strip()))

        if not found:
            return []
        lo, hi = found[0][0], found[-1][0]
        for col, label in unknown:
            if lo < col < hi:
                found.append((col, label))
                self.report.add_issue(
                    "info", "Добыча по месторождениям",
                    f"столбец {get_column_letter(col)} «{label}» отсутствует "
                    f"в справочнике месторождений — выведен как есть")

        out = []
        for col, name in sorted(found):
            value = _num(self.sheet.value(value_row, col))
            if value is not None:
                out.append((name, value))
        return out

    def _read_delivery(self) -> None:
        """Показатели сдачи нефти."""
        d = Delivery()
        for name in L.DELIVERY:
            if hasattr(d, name):
                setattr(d, name, self._num_at(name))
        self.report.delivery = d

    def _read_stock(self) -> None:
        """Таблица состояния фонда скважин: итог и месторождения."""
        r = self.report
        attrs = [k for k in L.STOCK.columns if k != "name"]

        def build(row: int) -> StockRow:
            """Строка таблицы фонда по номеру строки листа."""
            item = StockRow(name=_text(self._cell("stock", "name", row)))
            for attr in attrs:
                setattr(item, attr, _num(self._cell("stock", attr, row)))
            return item

        total_row = self.map.totals.get("stock")
        if total_row:
            r.stock_total = build(total_row)
        for row, _name in self._rows_of("stock", L.STOCK):
            item = build(row)
            if (item.total or 0) > 0:
                r.stock_by_field.append(item)
        if r.stock_total is None and r.stock_by_field:
            r.add_issue("warning", "Фонд скважин",
                        "строка итога по обществу не найдена")

    def _read_movement(self) -> None:
        """Запуски и остановки скважин за сутки."""
        r = self.report
        for row, well in self._rows_of("startups", L.STARTUPS):
            r.startups.append(WellRow(
                well=well,
                field_name=_text(self._cell("startups", "field", row)),
                qo=_num(self._cell("startups", "qo", row)),
                ql=_num(self._cell("startups", "ql", row)),
                wc=_num(self._cell("startups", "wc", row)),
                reason=_text(self._cell("startups", "reason", row)),
                note=_text(self._cell("startups", "note", row)),
            ))
        r.startups_total = self._summary_row(
            self.map.totals.get("startups"), self.map.col("startups", "well"))

        block = self.map.blocks.get("shutdowns")
        rows = {}
        for key, find in L.SHUTDOWN_ROWS.items():
            span = self.sheet.find(find.label, mode=find.mode, cols=find.cols,
                                   rows=block)
            rows[key] = span
        total = rows.get("total")
        r.shutdowns_total = self._summary_row(
            total.row if total else None, total.col if total else None)
        notes = []
        for key in ("svodka", "fact"):
            span = rows.get(key)
            if span is None:
                continue
            for col in range(span.col, min(span.col + 14, self.sheet.max_col) + 1):
                text = _text(self.sheet.raw(span.row, col))
                if len(text) > 12:
                    notes.append(text)
                    break
        r.shutdowns_note = " ".join(notes)

    def _summary_row(self, row: int | None, col: int | None) -> Totals:
        """Строка ИТОГО: подпись слева, дальше подряд количество и объёмы."""
        if row is None or col is None:
            return Totals()
        numbers = self.sheet.numbers_right(row, col - 1, limit=3, within=14)
        while len(numbers) < 3:
            numbers.append(None)
        return Totals(count=numbers[0], ql=numbers[1], qo=numbers[2])

    def _read_wells(self, table: str, spec: Table, block_name: str) -> None:
        """Список скважин простоя или бездействия.

        Args:
            table: Ключ таблицы в разрешённой схеме.
            spec: Описание таблицы из схемы.
            block_name: Название блока для текстов замечаний.
        """
        r = self.report
        items: list[WellRow] = []
        for row, well in self._rows_of(table, spec):
            item = WellRow(
                well=well,
                field_name=_text(self._cell(table, "field", row)),
                pump=_text(self._cell(table, "pump", row)),
                qo=_num(self._cell(table, "qo", row)),
                ql=_num(self._cell(table, "ql", row)),
                wc=_num(self._cell(table, "wc", row)),
                reason=_text(self._cell(table, "reason", row)),
                code=_text(self._cell(table, "code", row)),
                category=_text(self._cell(table, "category", row)),
                days=_text(self._cell(table, "days", row)),
                status=_text(self._cell(table, "status", row)),
                stop_date=_date(self._cell(table, "stop_date", row)),
                last_repair=_date(self._cell(table, "last_repair", row)),
                note=_text(self._cell(table, "note", row)),
            )
            for label, value in (("дата остановки", item.stop_date),
                                 ("дата последнего ремонта", item.last_repair)):
                if value and value.year > L.MAX_PLAUSIBLE_YEAR:
                    r.add_issue("warning", f"{block_name}, строка {row}",
                                f"скв. {well}: недостоверная {label} — "
                                f"{value:%d.%m.%Y}")
            items.append(item)

        totals = Totals(
            count=float(len(items)),
            qo=sum(i.qo or 0.0 for i in items),
            ql=sum(i.ql or 0.0 for i in items),
        )
        self._verify_count(table, block_name, len(items))

        if table == "idle":
            r.idle, r.idle_total = items, totals
        else:
            r.downtime, r.downtime_total = items, totals
            self._read_downtime_breakdown()

    def _verify_count(self, table: str, block_name: str, parsed: int) -> None:
        """Сверка с количеством скважин, объявленным в строке ИТОГО."""
        row = self.map.totals.get(table)
        col = self.map.col(table, "well")
        if row is None or col is None:
            return
        declared = _num(self.sheet.value(row, col))
        if declared is not None and abs(declared - parsed) > 0.5:
            self.report.add_issue(
                "warning", block_name,
                f"в строке ИТОГО указано {declared:.0f} скв., "
                f"прочитано {parsed}")

    def _read_downtime_breakdown(self) -> None:
        """Готовая разбивка простоя по причинам из строк перед ИТОГО.

        Причины не перечисляются заранее: их состав в рапорте меняют.
        В августе 2026 года «Сезонная эксплуатация» уступила место
        «Ожиданию СКО», и список известных подписей потерял бы 28,7 т
        из общего простоя. Поэтому читается то, что написано.

        Разбивка занимает строки, вплотную примыкающие к ИТОГО, и в
        каждой из них помещается по несколько групп «подпись —
        количество скважин — потери». Границей служит первая строка,
        в которой таких групп нет: выше идут сводки по трубопроводам
        с такими же по виду числами, и захватывать их нельзя.

        Искать подписи по всему блоку тоже нельзя: короткие вроде «КВД»
        встречаются в столбце «Причина остановки» у отдельных скважин.
        """
        r = self.report
        block = self.map.blocks.get("downtime")
        total_row = self.map.totals.get("downtime")
        if total_row is None:
            return
        limit = max(block[0] if block else 1, total_row - L.BREAKDOWN_DEPTH)
        for row in range(total_row - 1, limit - 1, -1):
            groups = self._breakdown_row(row)
            if not groups:
                break
            r.downtime_breakdown.extend(
                (name, count, qo) for name, count, qo in groups if count or qo)
        r.downtime_breakdown.sort(key=lambda x: x[2], reverse=True)

    def _breakdown_row(self, row: int) -> list[tuple[str, float, float]]:
        """Разобрать одну строку разбивки на группы «причина, скважины, потери».

        Args:
            row: Номер строки листа.

        Returns:
            Список троек в порядке следования по строке. Пустой список
            означает, что строка к разбивке не относится.
        """
        groups: list[tuple[str, float, float]] = []
        col = 1
        while col <= self.sheet.max_col:
            value = self.sheet.raw(row, col)
            if not isinstance(value, str) or not value.strip():
                col += 1
                continue
            span = self.sheet.span(row, col)
            numbers = self.sheet.numbers_right(row, span.last_col, limit=2,
                                               within=L.BREAKDOWN_SPAN)
            if len(numbers) == 2:
                groups.append((_text(value), numbers[0], numbers[1]))
            col = span.last_col + 1
        return groups

    def _read_vsp(self) -> None:
        """Внутрисменные потери: скважины и объём потерь."""
        r = self.report
        for row, well in self._rows_of("vsp", L.VSP):
            r.vsp.append(WellRow(
                well=well,
                field_name=_text(self._cell("vsp", "field", row)),
                pump=_text(self._cell("vsp", "pump", row)),
                qo=_num(self._cell("vsp", "loss_qo", row)),
                ql=_num(self._cell("vsp", "loss_ql", row)),
                wc=_num(self._cell("vsp", "wc", row)),
                reason=_text(self._cell("vsp", "reason", row)),
                code=_text(self._cell("vsp", "code", row)),
                category=_text(self._cell("vsp", "category", row)),
                days=_text(self._cell("vsp", "duration", row)),
                status=_text(self._cell("vsp", "stop_time", row)),
                note=_text(self._cell("vsp", "note", row)),
            ))
        r.vsp_total = Totals(
            count=float(len(r.vsp)),
            qo=sum(i.qo or 0.0 for i in r.vsp),
            ql=sum(i.ql or 0.0 for i in r.vsp),
        )

    def _ramp_rows(self, table: str, spec: Table) -> list[RampUpRow]:
        """Строки блока вывода на режим или контроля после ГТМ.

        Args:
            table: Ключ таблицы в разрешённой схеме.
            spec: Описание таблицы из схемы.

        Returns:
            Список скважин с плановыми и фактическими параметрами.
        """
        out = []
        for row, well in self._rows_of(table, spec):
            out.append(RampUpRow(
                field_name=_text(self._cell(table, "field", row)),
                well=well,
                plan_ql=_num(self._cell(table, "plan_ql", row)),
                plan_qo=_num(self._cell(table, "plan_qo", row)),
                fact_ql=_num(self._cell(table, "fact_ql", row)),
                fact_qo=_num(self._cell(table, "fact_qo", row)),
                fact_wc=_num(self._cell(table, "fact_wc", row)),
                pump=_text(self._cell(table, "pump", row)),
                gtm=_text(self._cell(table, "gtm", row)),
                days=_text(self._cell(table, "days", row)),
                note=_text(self._cell(table, "note", row)),
            ))
        return out

    def _read_ramp_up(self) -> None:
        """Вывод на режим: список скважин и суммарные показатели."""
        r = self.report
        r.ramp_up = self._ramp_rows("rampup", L.RAMP_UP)
        r.ramp_up_total = {
            "count": float(len(r.ramp_up)),
            "plan_ql": sum(i.plan_ql or 0.0 for i in r.ramp_up),
            "plan_qo": sum(i.plan_qo or 0.0 for i in r.ramp_up),
            "fact_ql": sum(i.fact_ql or 0.0 for i in r.ramp_up),
            "fact_qo": sum(i.fact_qo or 0.0 for i in r.ramp_up),
        }

    def _read_injection(self) -> None:
        """Поддержание пластового давления по объектам."""
        r = self.report
        attrs = [k for k in L.INJECTION.columns if k != "name"]

        def build(row: int) -> InjectionRow:
            """Строка таблицы ППД по номеру строки листа."""
            item = InjectionRow(name=_text(self._cell("injection", "name", row)))
            for attr in attrs:
                if hasattr(item, attr):
                    setattr(item, attr, _num(self._cell("injection", attr, row)))
            return item

        for row, _name in self._rows_of("injection", L.INJECTION):
            r.injection.append(build(row))
        total_row = self.map.totals.get("injection")
        r.injection_total = build(total_row) if total_row else None

    def _read_tanks(self) -> None:
        """Резервуарный парк: обе половины таблицы."""
        r = self.report
        for table in ("tanks_left", "tanks_right"):
            spec = L.TANKS_LEFT if table == "tanks_left" else L.TANKS_RIGHT
            for row, name in self._rows_of(table, spec):
                r.tanks.append(TankRow(
                    name=" ".join(name.split()),
                    level=_num(self._cell(table, "level", row)),
                    volume=_num(self._cell(table, "volume", row)),
                    delta=_num(self._cell(table, "delta", row)),
                    free=_num(self._cell(table, "free", row)),
                ))
        r.free_volume_total_m3 = self._num_at("free_volume_total")

    def _read_quality(self) -> None:
        """Качество товарной нефти по узлам учёта."""
        attrs = [k for k in L.OIL_QUALITY.columns if k != "node"]
        for row, node in self._rows_of("quality", L.OIL_QUALITY):
            item = QualityRow(node=node)
            for attr in attrs:
                setattr(item, attr, _num(self._cell("quality", attr, row)))
            self.report.quality.append(item)

    def _read_gas(self) -> None:
        """Баланс попутного газа и процент утилизации."""
        r = self.report
        for row, name in self._rows_of("gas", L.GAS):
            item = GasRow(
                name=name,
                day_plan=_num(self._cell("gas", "day_plan", row)),
                day_fact=_num(self._cell("gas", "day_fact", row)),
                day_dev=_num(self._cell("gas", "day_dev", row)),
                month_fact=_num(self._cell("gas", "month_fact", row)),
                year_fact=_num(self._cell("gas", "year_fact", row)),
            )
            if "%" in name:
                r.gas_utilization_pct = item.day_fact
            else:
                r.gas.append(item)

    def _read_power(self) -> None:
        """Энергетика: выработка, состав оборудования, остатки топлива."""
        p = PowerState()
        for attr in L.POWER:
            if hasattr(p, attr):
                setattr(p, attr, self._num_at(attr))
        for attr in ("total_sh_zh", "total_with_psp"):
            cell = self.map.cell(attr)
            if cell:
                setattr(p, f"{attr}_delta",
                        _num(self.sheet.value(cell[0],
                                              cell[1] + L.POWER_DELTA_DC)))

        cell = self.map.cell("diesel_stock")
        if cell:
            label_row, value_row = cell[0], cell[0] + 1
            for col in range(cell[1], min(cell[1] + 6, self.sheet.max_col) + 1):
                name = _text(self.sheet.raw(label_row, col))
                value = _num(self.sheet.raw(value_row, col))
                if name and value is not None:
                    p.diesel_stock.append((name, value))

        block = self.map.blocks.get("power")
        for find in L.POWER_NOTES:
            span = self.sheet.find(find.label, mode="starts", cols=find.cols,
                                   rows=block)
            if span is not None:
                text = _text(self.sheet.value(span.row, span.col))
                if len(text) > 40:
                    p.notes.append(text)
        self.report.power = p

    def _read_pigging(self) -> None:
        """Пуски и приёмы очистных устройств по участкам."""
        r = self.report
        rows = self.map.rows.get("pigging") or L.PIGGING.reference_rows
        line_col = self.map.col("pigging", "line")
        current = ""
        for row in range(rows[0], rows[1] + 1):
            if line_col:
                label = _text(self.sheet.raw(row, line_col))
                if label:
                    current = label
            launch = self._cell("pigging", "launch", row)
            receive = self._cell("pigging", "receive", row)
            if not isinstance(launch, datetime) and not isinstance(receive, datetime):
                continue
            travel = _num(self._cell("pigging", "travel_h", row))
            if travel is not None and not (0 <= travel <= L.PIGGING_MAX_HOURS):
                r.add_issue("warning", f"ОУ, строка {row}",
                            f"недостоверное время в пути: {_ru(travel)} ч")
                travel = None
            r.pigging.append(PiggingRow(
                line=current,
                section=_text(self._cell("pigging", "section", row)),
                launch=launch if isinstance(launch, datetime) else None,
                receive=receive if isinstance(receive, datetime) else None,
                travel_h=travel,
                note=_text(self._cell("pigging", "note", row)),
            ))

    def _read_drilling(self) -> None:
        """Скважины в бурении."""
        for row, well in self._rows_of("drilling", L.DRILLING):
            raw_depth = self._cell("drilling", "current_depth", row)
            self.report.drilling.append(DrillingRow(
                contractor=_text(self._cell("drilling", "contractor", row)),
                field_name=_text(self._cell("drilling", "field", row)),
                well=well,
                target_depth=_num(self._cell("drilling", "target_depth", row)),
                current_depth=_num(raw_depth),
                current_depth_text=_text(raw_depth),
                start_date=_date(self._cell("drilling", "start_date", row)),
            ))

    def _read_hse(self) -> None:
        """Происшествия и обращения в здравпункт."""
        hse = HSE()
        keys = list(L.HSE.columns)
        rows = self.map.rows.get("hse") or L.HSE.reference_rows
        for row in range(rows[0], rows[1] + 1):
            values = {k: _text(self._cell("hse", k, row)) for k in keys}
            if any(v and not _is_empty(v) for v in values.values()):
                hse.incidents.append(values)
        raw = self._text_at("hse_medical")
        hse.medical_raw = raw
        # персональные данные в презентацию не выносим, считаем только обращения
        hse.medical_count = len(re.findall(r"обрат", raw, flags=re.IGNORECASE))
        self.report.hse = hse

    # --- происхождение показателей -------------------------------------------
    def _collect_sources(self) -> None:
        """Записать, откуда взялся каждый показатель презентации.

        Прямые значения берутся из уже разрешённой схемы, поэтому адреса
        ячеек не дублируются в коде. Суммы и расчёты описываются явно.
        """
        ledger = self.report.ledger
        specs = [
            ("Шапка", L.HEADER),
            ("Добыча", L.PRODUCTION),
            ("Сдача", L.DELIVERY),
            ("Энергетика", L.POWER),
        ]
        for section, spec in specs:
            for key, item in spec.items():
                cell = self.map.cell(key)
                if cell is None:
                    continue
                label = L.METRIC_LABELS.get(key) or (
                    item.find.label if hasattr(item, "find")
                    else f"{item.row.label} → {item.col.label}")
                self._register_cell(ledger, section, label, cell)

        for key, section, label in (
            ("free_volume_total", "Сдача", "Свободный объём по всем объектам"),
            ("gas_consumers_total", "Газ", "Расход газа, всего"),
            ("diesel_left", "Энергетика", "Остаток дизельного топлива"),
        ):
            cell = self.map.cell(key)
            if cell:
                self._register_cell(ledger, section, label, cell)

        self._sources_for_stock(ledger)
        self._sources_for_lists(ledger)
        self._sources_for_derived(ledger)

    def _register_cell(self, ledger, section: str, label: str,
                       cell: tuple[int, int]) -> None:
        """Записать источник, указав адрес по началу объединённого диапазона.

        Значение объединённой ячейки хранится в её левом верхнем углу,
        поэтому в отчёте должен стоять именно этот адрес — иначе проверяющий
        откроет пустую ячейку.
        """
        span = self.sheet.span(*cell)
        ledger.cell(section, label, self.sheet.value(*cell),
                    span.row, span.col)

    def _sources_for_stock(self, ledger) -> None:
        """Источники показателей таблицы фонда скважин."""
        row = self.map.totals.get("stock")
        stock = self.report.stock_total
        if row is None or stock is None:
            return
        for attr, label in (
            ("in_work", "Фонд: в работе"),
            ("in_work_delta", "Фонд: изменение за сутки"),
            ("in_downtime", "Фонд: в текущем простое, скв."),
            ("downtime_rate", "Фонд: текущий простой, т/сут"),
            ("in_idle", "Фонд: в бездействии, скв."),
            ("idle_rate", "Фонд: бездействие, т/сут"),
        ):
            col = self.map.col("stock", attr)
            if col:
                self._register_cell(ledger, "Фонд скважин", label, (row, col))

    def _sources_for_lists(self, ledger) -> None:
        """Источники сумм по спискам скважин."""
        r = self.report
        stock = r.stock_total
        lists = (
            ("Текущий простой", "downtime", r.downtime, r.downtime_total,
             stock.downtime_rate if stock else None, "qo"),
            ("Бездействие", "idle", r.idle, r.idle_total,
             stock.idle_rate if stock else None, "qo"),
            ("Внутрисменные потери", "vsp", r.vsp, r.vsp_total, None,
             "loss_qo"),
        )
        for section, table, items, totals, declared, column in lists:
            rows = self.map.rows.get(table)
            col = self.map.col(table, column)
            if rows is None or col is None:
                continue
            check = ""
            if declared is not None:
                check = f"таблицей фонда ({_ru(declared, 3)})"
            ledger.total(section, f"{section}: потери, т/сут", totals.qo,
                         col, rows, len(items), check)
            ledger.count(section, f"{section}: количество скважин",
                         len(items),
                         f"число заполненных строк {rows[0]}–{rows[1]}")

    def _sources_for_derived(self, ledger) -> None:
        """Источники величин, посчитанных программой."""
        r = self.report
        parts = (
            ("текущий простой", r.downtime_total.qo),
            ("бездействие", r.idle_total.qo),
            ("внутрисменные", r.vsp_total.qo),
        )
        formula = " + ".join(f"{_ru(value or 0, 3)} ({name})"
                             for name, value in parts)
        ledger.derived("Потери добычи", "Потери добычи, всего за сутки",
                       r.total_losses, formula)

        if r.losses_share is not None:
            ledger.derived(
                "Потери добычи", "Доля потерь в суточной добыче",
                f"{_ru(r.losses_share, 1)} %",
                f"{_ru(r.total_losses, 3)} / "
                f"{_ru(r.production.closed_day or 0, 3)} × 100")

        p = r.production
        if p.potential_com is not None and p.potential_limited is not None:
            ledger.derived(
                "Добыча", "Потери потенциала от техограничений",
                p.potential_limited - p.potential_com,
                f"{_ru(p.potential_limited, 3)} (с учётом ограничений) − "
                f"{_ru(p.potential_com, 3)} (потенциал)")
        if p.fact_com_day is not None and p.potential_limited is not None:
            ledger.derived(
                "Добыча", "Отклонение факта от ограниченного потенциала",
                p.fact_com_day - p.potential_limited,
                f"{_ru(p.fact_com_day, 3)} (факт ЦОМ) − "
                f"{_ru(p.potential_limited, 3)}")

        if r.ramp_up_total:
            ledger.derived(
                "Вывод на режим", "Вывод на режим: факт, т/сут",
                r.ramp_up_total.get("fact_qo"),
                f"сумма по {len(r.ramp_up)} скважинам блока «Вывод на режим»")

        ledger.count("ПБиОТ", "Обращений в здравпункт", r.hse.medical_count,
                     "число упоминаний «обрат…» в тексте блока")
        ledger.count("Бурение", "Скважин в бурении", len(r.drilling),
                     "число заполненных строк блока «Буровой подрядчик»")


def span_text(sheet: Sheet, span) -> str:
    """Текст левой верхней ячейки диапазона."""
    return sheet.text(span.row, span.col)


def read_report(path: str) -> DailyReport:
    """Прочитать суточный рапорт.

    Args:
        path: Путь к файлу .xlsx.

    Returns:
        Заполненную модель суточного рапорта.
    """
    return ReportReader(path).read()


def layout_report(path: str) -> tuple[Resolved, Sheet]:
    """Только привязка схемы — для режима --check."""
    sheet = open_sheet(path)
    return resolve(sheet), sheet
