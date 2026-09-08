# -*- coding: utf-8 -*-
"""Происхождение каждого показателя, попадающего в презентацию.

Требование к докладу: любую цифру на слайде можно проверить. Поэтому
программа сопровождает каждый показатель ссылкой на источник — адресом
ячейки рапорта, описанием суммируемого диапазона или формулой расчёта.

Расчётные величины на слайдах помечаются словом «расч.», а полный перечень
источников выводится ключом ``--trace`` и кнопкой «Откуда цифры» в окне.
"""
from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter

#: Вид происхождения показателя.
CELL = "cell"          # прямое значение ячейки рапорта
SUM = "sum"            # сумма диапазона строк
DERIVED = "derived"    # арифметика над другими показателями
COUNT = "count"        # количество строк или совпадений

KIND_TITLES = {
    CELL: "Из ячейки рапорта",
    SUM: "Сумма по строкам рапорта",
    DERIVED: "Расчёт по другим показателям",
    COUNT: "Количество записей",
}

#: Пометка расчётных величин на слайдах.
DERIVED_MARK = "расч."


def address(row: int, col: int) -> str:
    """Координаты ячейки в привычном виде: «K28»."""
    return f"{get_column_letter(col)}{row}"


@dataclass(frozen=True)
class Source:
    """Одна строка отчёта о происхождении."""
    section: str        # раздел рапорта или слайд
    label: str          # как показатель называется
    value: str          # отформатированное значение
    origin: str         # ячейка, диапазон или формула
    kind: str = CELL

    @property
    def is_derived(self) -> bool:
        """Получена ли величина расчётом, а не взята из ячейки."""
        return self.kind in (DERIVED, COUNT)


@dataclass
class Ledger:
    """Накопитель сведений о происхождении показателей."""
    items: list[Source] = field(default_factory=list)

    def cell(self, section: str, label: str, value, row: int, col: int) -> None:
        """Записать показатель, взятый напрямую из ячейки рапорта.

        Args:
            section: Раздел рапорта.
            label: Название показателя.
            value: Прочитанное значение.
            row: Строка ячейки.
            col: Столбец ячейки.
        """
        self.items.append(Source(section, label, _fmt(value),
                                 f"ячейка {address(row, col)}", CELL))

    def total(self, section: str, label: str, value, col: int,
              rows: tuple[int, int], count: int, check: str = "") -> None:
        """Записать показатель, полученный суммированием строк.

        Args:
            section: Раздел рапорта.
            label: Название показателя.
            value: Полученная сумма.
            col: Суммируемый столбец.
            rows: Границы диапазона строк.
            count: Сколько значений просуммировано.
            check: С чем результат сверен, если сверка выполнялась.
        """
        origin = (f"сумма столбца {get_column_letter(col)}, "
                  f"строки {rows[0]}–{rows[1]} ({count} знач.)")
        if check:
            origin += f"; сверено с {check}"
        self.items.append(Source(section, label, _fmt(value), origin, SUM))

    def derived(self, section: str, label: str, value,
                formula: str) -> None:
        """Записать величину, посчитанную из других показателей.

        Args:
            section: Раздел рапорта.
            label: Название показателя.
            value: Результат расчёта.
            formula: Формула в виде текста — попадёт в отчёт как есть.
        """
        self.items.append(Source(section, label, _fmt(value), formula, DERIVED))

    def count(self, section: str, label: str, value, origin: str) -> None:
        """Записать количество записей.

        Args:
            section: Раздел рапорта.
            label: Название показателя.
            value: Полученное количество.
            origin: Пояснение, что именно подсчитано.
        """
        self.items.append(Source(section, label, _fmt(value), origin, COUNT))

    # --- выборки -------------------------------------------------------------
    @property
    def derived_items(self) -> list[Source]:
        """Только расчётные величины и подсчёты записей."""
        return [item for item in self.items if item.is_derived]


def _fmt(value) -> str:
    """Значение показателя в виде текста для отчёта."""
    if value is None:
        return "—"
    if isinstance(value, float):
        text = f"{value:,.3f}".rstrip("0").rstrip(".")
        return text.replace(",", " ").replace(".", ",")
    if isinstance(value, int):
        return f"{value:,}".replace(",", " ")
    return str(value)


def render(ledger: Ledger, only_derived: bool = False) -> str:
    """Отчёт о происхождении показателей в виде таблицы."""
    items = ledger.derived_items if only_derived else ledger.items
    if not items:
        return "Показатели не зарегистрированы."

    title = ("РАСЧЁТНЫЕ ВЕЛИЧИНЫ" if only_derived
             else "ПРОИСХОЖДЕНИЕ ПОКАЗАТЕЛЕЙ")
    lines = [title, "=" * len(title), ""]

    order = [DERIVED, COUNT, SUM, CELL]
    for kind in order:
        chosen = [item for item in items if item.kind == kind]
        if not chosen:
            continue
        lines.append(f"--- {KIND_TITLES[kind]} ({len(chosen)}) ---")
        width = min(52, max(len(i.label) for i in chosen) + 2)
        for item in chosen:
            label = item.label if len(item.label) <= width else \
                item.label[:width - 1] + "…"
            lines.append(f"  {label:<{width}} {item.value:>14}   "
                         f"{item.origin}")
        lines.append("")

    lines.append("Пометка «расч.» на слайде означает, что величина получена "
                 "арифметически и в самом рапорте отсутствует.")
    return "\n".join(lines)
